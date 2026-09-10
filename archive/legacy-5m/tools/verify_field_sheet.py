from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


path = Path("/workspace/sites/invasive-plant-transect-system/field-sheet/Invasive_Plant_Transect_Field_Sheet.docx")
doc = Document(path)

assert len(doc.sections) == 2, f"Expected 2 Word sections, found {len(doc.sections)}"
assert all(section.orientation == WD_ORIENT.LANDSCAPE for section in doc.sections)

survey_tables = [table for table in doc.tables if len(table.rows) == 17 and len(table.columns) == 11]
assert len(survey_tables) == 2, f"Expected two 17x11 survey tables, found {len(survey_tables)}"

expected_pages = [
    [f"{index}-{index + 1}" for index in range(0, 15)],
    [f"{index}-{index + 1}" for index in range(15, 30)],
]
for table, expected in zip(survey_tables, expected_pages):
    actual = [table.cell(row_index, 5).text.strip() for row_index in range(2, 17)]
    assert actual == expected, (actual, expected)
    header_order = [table.cell(1, column).text.strip() for column in range(11)]
    assert header_order == ["4-5 m", "3-4 m", "2-3 m", "1-2 m", "0-1 m", "SEGMENT", "0-1 m", "1-2 m", "2-3 m", "3-4 m", "4-5 m"]
    grid_widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
    tbl_width = int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")))
    assert sum(grid_widths) == tbl_width == 14832
    for row in table.rows:
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None

document_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
document_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
for required in [
    "Date:", "Observer(s):", "Site:", "Trail:", "Transect #:", "Start time:",
    "End time:", "Start GPS:", "End GPS:", "General notes:",
    "LEFT - FACING DIRECTION OF TRAVEL", "RIGHT - FACING DIRECTION OF TRAVEL",
    "0 = surveyed/no target", "NS = not surveyed", "blank = incomplete",
]:
    assert required in document_text, f"Missing text: {required}"
assert "species-code legend" not in document_text.lower()

for paragraph in list(doc.paragraphs) + [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]:
    for run in paragraph.runs:
        if run.text and run.font.name:
            assert run.font.name == "Times New Roman", f"Unexpected run font {run.font.name}: {run.text!r}"

with ZipFile(path) as archive:
    names = set(archive.namelist())
    assert "word/document.xml" in names
    assert "[Content_Types].xml" in names

print({"sections": 2, "survey_tables": 2, "segments": 30, "cells": 300, "font": "Times New Roman", "geometry": "pass"})
