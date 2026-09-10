from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/workspace/sites/invasive-plant-transect-system/field-sheet/Invasive_Plant_Transect_Field_Sheet.docx")

# compact_reference_guide preset with named field-form overrides:
# Letter landscape; 0.35-inch margins; Times New Roman throughout; 10.3-inch fixed tables;
# restrained grayscale/pale fills; compact, writable form spacing.
PAGE_WIDTH_IN = 11.0
PAGE_HEIGHT_IN = 8.5
SIDE_MARGIN_IN = 0.35
VERTICAL_MARGIN_IN = 0.18
CONTENT_WIDTH_IN = PAGE_WIDTH_IN - (2 * SIDE_MARGIN_IN)
CONTENT_WIDTH_DXA = round(CONTENT_WIDTH_IN * 1440)
TABLE_INDENT_DXA = 0
CELL_MARGIN_DXA = {"top": 35, "bottom": 35, "start": 45, "end": 45}
FONT = "Times New Roman"
INK = "000000"
MID_GRAY = "D9D9D9"
LIGHT_GRAY = "F1F1F1"
LEFT_FILL = "E3EAF2"
RIGHT_FILL = "EEE5E1"
WHITE = "FFFFFF"


def set_run_font(run, size=8, bold=False, italic=False, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, before=0, after=0, line=1.0, alignment=None, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    if alignment is not None:
        paragraph.alignment = alignment


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margins=CELL_MARGIN_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in margins.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            set_cell_width(cell, width)
            set_cell_margins(cell)


def set_table_borders(table, size=8, color="4A4A4A", inside_size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(inside_size if edge.startswith("inside") else size))
        node.set(qn("w:color"), color)


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_row_min_height(row, inches):
    row.height = Inches(inches)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    prevent_row_split(row)


def clear_cell(cell):
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    return paragraph


def write_cell(cell, text, size=8, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, alignment=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade_cell(cell, fill)
    return paragraph


def add_label_line(cell, label, blank="", size=8, notes=False):
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=size, bold=True)
    value_run = paragraph.add_run(blank)
    set_run_font(value_run, size=size)
    if notes:
        spacer = cell.add_paragraph()
        set_paragraph(spacer, before=0, after=0, line=1.0)
        set_run_font(spacer.add_run(""), size=size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_section_geometry(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(VERTICAL_MARGIN_IN)
    section.bottom_margin = Inches(VERTICAL_MARGIN_IN)
    section.left_margin = Inches(SIDE_MARGIN_IN)
    section.right_margin = Inches(SIDE_MARGIN_IN)
    section.header_distance = Inches(0.08)
    section.footer_distance = Inches(0.08)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=7)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def set_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    set_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(paragraph.add_run("Invasive Plant Transect Field Sheet | Page "), size=7)
    add_page_number_field(paragraph)
    set_run_font(paragraph.add_run(" of 2"), size=7)


def add_title(doc, page_label):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, before=0, after=0, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    set_run_font(paragraph.add_run("INVASIVE PLANT TRANSECT FIELD SHEET"), size=12.5, bold=True)
    subtitle = doc.add_paragraph()
    set_paragraph(subtitle, before=0, after=2, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    set_run_font(subtitle.add_run(page_label), size=8.5, bold=True)


def add_metadata_page_one(doc):
    table = doc.add_table(rows=4, cols=6)
    widths = [2472] * 6
    set_table_geometry(table, widths)
    set_table_borders(table, size=6, inside_size=4)

    cells = table.rows[0].cells
    date = cells[0].merge(cells[1])
    site = cells[2].merge(cells[3])
    transect = cells[4].merge(cells[5])
    add_label_line(date, "Date")
    add_label_line(site, "Site")
    add_label_line(transect, "Transect #")
    set_row_min_height(table.rows[0], 0.30)

    cells = table.rows[1].cells
    observers = cells[0].merge(cells[3])
    trail = cells[4].merge(cells[5])
    add_label_line(observers, "Observer(s)")
    add_label_line(trail, "Trail")
    set_row_min_height(table.rows[1], 0.30)

    cells = table.rows[2].cells
    start_time = cells[0]
    end_time = cells[1]
    start_gps = cells[2].merge(cells[3])
    end_gps = cells[4].merge(cells[5])
    add_label_line(start_time, "Start time", size=7.5)
    add_label_line(end_time, "End time", size=7.5)
    add_label_line(start_gps, "Start GPS", size=7.5)
    add_label_line(end_gps, "End GPS", size=7.5)
    set_row_min_height(table.rows[2], 0.30)

    notes = table.rows[3].cells[0].merge(table.rows[3].cells[5])
    add_label_line(notes, "General notes", size=7.5, notes=True)
    set_row_min_height(table.rows[3], 0.43)


def add_metadata_page_two(doc):
    table = doc.add_table(rows=1, cols=3)
    widths = [5760, 3600, 5472]
    set_table_geometry(table, widths)
    set_table_borders(table, size=6, inside_size=4)
    add_label_line(table.cell(0, 0), "Site")
    add_label_line(table.cell(0, 1), "Transect #")
    add_label_line(table.cell(0, 2), "Date")
    set_row_min_height(table.rows[0], 0.33)


def add_instructions(doc):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, size=6, inside_size=4)
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GRAY)
    paragraph = clear_cell(cell)
    set_paragraph(paragraph, before=0, after=0, line=1.0)
    set_run_font(paragraph.add_run("HOW TO RECORD: "), size=7.2, bold=True)
    set_run_font(paragraph.add_run(
        "Each cell is 1 m along trail x 1 m outward from the trail centerline. "
        "Assign a plant by its rooted location. LEFT/RIGHT always face from transect start toward end. "
        "Enter each species code once per cell; 0 = surveyed/no target; NS = not surveyed; blank = incomplete."
    ), size=7.2)
    set_row_min_height(table.rows[0], 0.35)


def add_survey_table(doc, start_segment, body_row_height):
    table = doc.add_table(rows=17, cols=11)
    band_width = 1368
    center_width = CONTENT_WIDTH_DXA - (10 * band_width)
    widths = [band_width] * 5 + [center_width] + [band_width] * 5
    set_table_geometry(table, widths)
    set_table_borders(table, size=8, inside_size=5)

    # Super-header: far-to-near on left, near-to-far on right.
    top = table.rows[0].cells
    left = top[0].merge(top[4])
    trail = top[5]
    right = top[6].merge(top[10])
    write_cell(left, "LEFT - FACING DIRECTION OF TRAVEL", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=LEFT_FILL)
    write_cell(trail, "TRAIL", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=MID_GRAY)
    write_cell(right, "RIGHT - FACING DIRECTION OF TRAVEL", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=RIGHT_FILL)
    set_row_min_height(table.rows[0], 0.22)
    mark_repeat_header(table.rows[0])

    headers = ["4-5 m", "3-4 m", "2-3 m", "1-2 m", "0-1 m", "SEGMENT", "0-1 m", "1-2 m", "2-3 m", "3-4 m", "4-5 m"]
    for index, text in enumerate(headers):
        fill = MID_GRAY if index == 5 else LIGHT_GRAY
        write_cell(table.cell(1, index), text, size=7.2, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=fill)
    set_row_min_height(table.rows[1], 0.24)
    mark_repeat_header(table.rows[1])

    for row_index in range(15):
        segment_start = start_segment + row_index
        row = table.rows[row_index + 2]
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if column_index == 5:
                write_cell(
                    cell,
                    f"{segment_start}-{segment_start + 1}",
                    size=7.5,
                    bold=True,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    fill="F7F7F7",
                )
            else:
                write_cell(cell, "", size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_row_min_height(row, body_row_height)
    return table


def configure_styles(doc):
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def build():
    doc = Document()
    configure_styles(doc)
    core = doc.core_properties
    core.title = "Invasive Plant Transect Field Sheet"
    core.subject = "30-meter, 300-cell invasive plant presence survey"
    core.author = ""
    core.keywords = "invasive plants; transect; field sheet"

    first = doc.sections[0]
    set_section_geometry(first)
    set_footer(first)
    add_title(doc, "PAGE 1 - TRAIL DISTANCE 0-15 m (segments 0-1 through 14-15 m)")
    add_metadata_page_one(doc)
    add_instructions(doc)
    add_survey_table(doc, 0, 0.27)

    second = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(second)
    second.header.is_linked_to_previous = True
    second.footer.is_linked_to_previous = True
    add_title(doc, "PAGE 2 - TRAIL DISTANCE 15-30 m (segments 15-16 through 29-30 m)")
    add_metadata_page_two(doc)
    add_instructions(doc)
    add_survey_table(doc, 15, 0.335)

    # Keep compatibility settings predictable across Word and LibreOffice.
    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    do_not_autofit = OxmlElement("w:doNotAutofitConstrainedTables")
    compat.append(do_not_autofit)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
